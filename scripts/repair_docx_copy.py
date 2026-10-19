
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from docx import Document


ROOT = Path(__file__).resolve().parents[1]

source = ROOT / "books" / (
    "17 崔璀：分析力发挥（下），"
    "不擅交际也能实现高效合作.docx"
)

output_dir = ROOT / "data" / "repaired"
output_dir.mkdir(parents=True, exist_ok=True)

target = output_dir / "17_分析力发挥下_repaired.docx"

REL_FILE = "word/_rels/document.xml.rels"
DOC_FILE = "word/document.xml"

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}

with ZipFile(source, "r") as src:
    relationships = etree.fromstring(src.read(REL_FILE))
    document = etree.fromstring(src.read(DOC_FILE))

    # 1. 找到确认损坏的图片关系
    matches = [
        rel for rel in relationships
        if rel.get("Id") == "rId6"
        and rel.get("Target") == "../NULL"
        and rel.get("Type", "").endswith("/image")
    ]

    if len(matches) != 1:
        raise RuntimeError("图片关系与诊断结果不一致")

    # 2. 查找正文中的图片引用
    blips = document.xpath(
        ".//a:blip[@r:embed='rId6']",
        namespaces=NS
    )

    if len(blips) != 1:
        raise RuntimeError("图片引用数量异常")

    # 3. 定位包含损坏图片的绘图元素
    drawings = blips[0].xpath(
        "ancestor::w:drawing[1]",
        namespaces=NS
    )

    if len(drawings) != 1:
        raise RuntimeError("无法安全定位绘图元素")

    drawing = drawings[0]

    # 如果绘图同时包含其他图片或文字，不自动删除
    if len(drawing.xpath(".//a:blip", namespaces=NS)) != 1:
        raise RuntimeError("绘图包含其他图片，停止修复")

    if drawing.xpath(".//w:t", namespaces=NS):
        raise RuntimeError("绘图包含文字，停止修复")

    # 4. 在内存中移除失效绘图和关系
    drawing.getparent().remove(drawing)
    relationships.remove(matches[0])

    # 5. 将修改结果保存到新 DOCX
    with ZipFile(target, "w") as dst:
        for info in src.infolist():

            if info.filename == REL_FILE:
                content = etree.tostring(
                    relationships,
                    encoding="UTF-8",
                    xml_declaration=True
                )

            elif info.filename == DOC_FILE:
                content = etree.tostring(
                    document,
                    encoding="UTF-8",
                    xml_declaration=True
                )

            else:
                content = src.read(info.filename)

            dst.writestr(info, content)

print("副本已保存：", target)

# 6. 检查是否能正常打开
doc = Document(target)

print("普通段落数：", len(doc.paragraphs))
print("表格数：", len(doc.tables))
print("副本读取成功")
