
import unittest
import tempfile
from pathlib import Path

from docx import Document

from app.ingestion.docx_reader import read_docx


class TestDOCXReader(unittest.TestCase):

    def test_read_docx(self):

        with tempfile.TemporaryDirectory() as temp_dir:

            path = Path(temp_dir) / "demo.docx"

            # 创建测试 Word
            doc = Document()
            doc.add_paragraph("分析力优势")
            doc.add_paragraph("")
            doc.add_paragraph("喜欢分析问题和寻找原因。")
            doc.save(path)

            # 调用读取器
            records = read_docx(path)

            # 验证结果
            self.assertEqual(len(records), 2)

            self.assertEqual(
                records[0].text,
                "分析力优势"
            )

            self.assertEqual(
                records[1].location,
                "第3段"
            )

            self.assertEqual(
                records[0].method,
                "docx_text"
            )

    def test_missing_file(self):

        with tempfile.TemporaryDirectory() as temp_dir:

            path = Path(temp_dir) / "missing.docx"

            with self.assertRaises(FileNotFoundError):
                read_docx(path)
    
    def test_paragraph_and_table_order(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "table_demo.docx"

            doc = Document()

            doc.add_paragraph("学习力的三种表现")

            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "类型"
            table.cell(0, 1).text = "表现"
            table.cell(1, 0).text = "行动学习者"
            table.cell(1, 1).text = "做中学"

            doc.add_paragraph("选择适合自己的学习方式")

            doc.save(path)

            records = read_docx(path)

            self.assertEqual(len(records), 4)

            self.assertEqual(
                records[0].text,
                "学习力的三种表现"
            )

            self.assertEqual(
                records[1].text,
                "类型 | 表现"
            )

            self.assertEqual(
                records[2].text,
                "行动学习者 | 做中学"
            )

            self.assertEqual(
                records[2].location,
                "第1表第2行"
            )

            self.assertEqual(
                records[3].text,
                "选择适合自己的学习方式"
            )



if __name__ == "__main__":
    unittest.main()
