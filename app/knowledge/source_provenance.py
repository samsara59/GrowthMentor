
"""原始教材来源解析与文本匹配。不会修改索引或执行OCR。"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict, dataclass
from functools import lru_cache
from pathlib import Path


SUPPORTED = {
    ".pdf", ".docx", ".txt", ".md",
}

IMAGES = {
    ".png", ".jpg", ".jpeg",
    ".tif", ".tiff", ".bmp",
}

INVENTORY = SUPPORTED | IMAGES | {
    ".doc", ".pptx", ".ppt",
    ".rtf", ".xlsx", ".xls", ".epub",
}


# ============================================
# 1. 文本比对只消除空白差异
# ============================================

def compact(text: str) -> str:

    return "".join(
        ch for ch in text
        if not ch.isspace()
    )


# ============================================
# 2. 原文核验结果
# ============================================

@dataclass(frozen=True)
class Verification:

    source: str
    status: str
    resolved_path: str = ""
    located_at: str = ""
    explanation: str = ""

    def as_dict(self) -> dict:
        return asdict(self)


# ============================================
# 3. 建立原始教材目录
# ============================================

class SourceCatalog:

    def __init__(
        self,
        books_dir: str | Path,
    ):

        self.root = (
            Path(books_dir).resolve()
        )

        if not self.root.is_dir():
            raise FileNotFoundError(
                f"教材目录不存在：{self.root}"
            )

        self.paths = []

        for p in sorted(
            self.root.rglob("*")
        ):

            if not p.is_file():
                continue

            if p.suffix.lower() not in INVENTORY:
                continue

            safe = p.resolve()

            # 不读取目录外的符号链接目标
            if not safe.is_relative_to(self.root):
                continue

            self.paths.append(safe)

        self.by_relative = {}
        self.by_name = defaultdict(list)

        for p in self.paths:

            relative = p.relative_to(
                self.root
            ).as_posix()

            self.by_relative[
                relative.casefold()
            ] = p

            self.by_name[
                p.name.casefold()
            ].append(p)

    def display(self, p: Path) -> str:

        return p.relative_to(
            self.root
        ).as_posix()

    # ========================================
    # 4. 根据元数据寻找真实文件
    # ========================================

    def resolve(
        self,
        raw_source: str | None,
    ) -> tuple[str, Path | None]:

        if (
            not isinstance(raw_source, str)
            or not raw_source.strip()
        ):

            return "source_missing", None

        name = raw_source.strip().replace(
            "\\", "/"
        )

        folded = name.casefold()

        # 兼容原始绝对路径中的 books/ 后缀
        if "/books/" in folded:

            name = name[
                folded.rfind("/books/")
                + len("/books/"):
            ]

        name = name.removeprefix(
            "./"
        ).lstrip("/")

        if name.casefold().startswith("books/"):
            name = name[len("books/"):]

        # 优先用相对于books的完整路径
        relative = self.by_relative.get(
            name.casefold()
        )

        if relative is not None:
            return "resolved", relative

        # 再尝试文件名，必须唯一
        candidates = self.by_name.get(
            name.split("/")[-1].casefold(),
            [],
        )

        if len(candidates) == 1:
            return "resolved", candidates[0]

        if len(candidates) > 1:
            return "ambiguous_source", None

        return "source_not_found", None


# ============================================
# 5. 从真实原始文件提取可比对文本
# ============================================

def _extract_original(
    path: Path,
) -> tuple[str, str, list[str]]:

    ext = path.suffix.lower()

    if ext in IMAGES:
        return "", "needs_ocr", []

    if ext not in SUPPORTED:
        return "", "unsupported_format", []

    try:

        # PDF：PyMuPDF
        if ext == ".pdf":

            import fitz

            with fitz.open(str(path)) as pdf:

                pages = [
                    page.get_text("text")
                    for page in pdf
                ]

            if not any(
                p.strip() for p in pages
            ):
                return "", "needs_ocr", []

            return (
                "\n".join(pages),
                "extracted",
                pages,
            )

        # DOCX：正文与表格
        if ext == ".docx":

            from docx import Document

            doc = Document(str(path))

            segments = [
                p.text
                for p in doc.paragraphs
            ]

            for table in doc.tables:

                for row in table.rows:

                    for cell in row.cells:

                        segments.append(
                            cell.text
                        )

            return (
                "\n".join(segments),
                "extracted",
                segments,
            )

        # 纯文本
        try:

            text = path.read_text(
                encoding="utf-8-sig"
            )

        except UnicodeDecodeError:

            text = path.read_text(
                encoding="gb18030"
            )

        return text, "extracted", [text]

    except Exception as exc:

        return (
            "",
            f"read_error:{type(exc).__name__}",
            [],
        )


@lru_cache(maxsize=8)
def _cached_extract(
    path: str,
) -> tuple[str, str, list[str]]:

    return _extract_original(
        Path(path)
    )


# ============================================
# 6. 对比一条检索片段与原文件
# ============================================

def verify_excerpt(
    catalog: SourceCatalog,
    source: str,
    excerpt: str,
) -> Verification:

    resolution, path = catalog.resolve(
        source
    )

    if path is None:

        return Verification(
            source,
            resolution,
            explanation=(
                "无法唯一定位原始教材"
            ),
        )

    relative = catalog.display(
        path
    )

    # 过短文本容易发生偶然匹配
    if (
        not isinstance(excerpt, str)
        or len(compact(excerpt)) < 30
    ):

        return Verification(
            source,
            "text_too_short",
            relative,
            explanation=(
                "候选片段过短，避免偶然匹配"
            ),
        )

    original, status, sections = (
        _cached_extract(str(path))
    )

    if status != "extracted":

        return Verification(
            source,
            status,
            relative,
            explanation=(
                "无法用当前文本解析方式核验，"
                "不得标记为原文已匹配"
            ),
        )

    needle = compact(excerpt)

    # 必须是整个片段都能在原文中匹配
    if needle not in compact(original):

        return Verification(
            source,
            "text_not_located",
            relative,
            explanation=(
                "文本清洗、OCR或跨页可能"
                "造成差异，不能直接认定引用虚构"
            ),
        )

    location = "document"

    if path.suffix.lower() == ".pdf":

        for index, page in enumerate(
            sections,
            start=1,
        ):

            if needle in compact(page):

                location = f"page={index}"

                break

        else:

            location = (
                "cross_page_or_unlocated"
            )

    return Verification(
        source,
        "matched_original_text",
        relative,
        location,
        (
            "原文字符序列匹配成功；"
            "尚不能证明该片段支持模型推断"
        ),
    )


# ============================================
# 7. 识别可能需要OCR或新解析器的文件
# ============================================

def suspected_ocr(
    catalog: SourceCatalog,
    candidates: list[Path],
    max_sample_pages: int = 4,
) -> list[dict]:

    result = []

    for path in candidates:

        ext = path.suffix.lower()

        if ext in IMAGES:

            result.append({
                "file": catalog.display(path),
                "reason": "image_requires_ocr",
            })

        elif ext == ".pdf":

            try:

                import fitz

                with fitz.open(
                    str(path)
                ) as pdf:

                    count = len(pdf)

                    # 兼容只有一页的PDF
                    indices = sorted(
                        i for i in {
                            0,
                            1,
                            count // 2,
                            count - 1,
                        }
                        if 0 <= i < count
                    )[:max_sample_pages]

                    texts = [
                        pdf[i].get_text("text")
                        for i in indices
                    ]

                if count and not any(
                    t.strip() for t in texts
                ):

                    result.append({
                        "file": catalog.display(path),
                        "reason": (
                            "sampled_pages_no_extractable_text"
                        ),
                        "pages_sampled": [
                            i + 1
                            for i in indices
                        ],
                    })

            except Exception as exc:

                result.append({
                    "file": catalog.display(path),
                    "reason": (
                        "pdf_probe_failed:"
                        + type(exc).__name__
                    ),
                })

        elif ext not in SUPPORTED:

            result.append({
                "file": catalog.display(path),
                "reason": (
                    "unsupported_extraction_format"
                ),
            })

    return result
